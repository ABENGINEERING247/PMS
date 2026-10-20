import os, json, io, datetime
import streamlit as st
import pandas as pd
from openai import OpenAI
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

st.set_page_config(page_title="PMS | Project Management System", page_icon="📊", layout="wide")
NAVY, GOLD, WHITE = "#12345A", "#D6A62E", "#FFFFFF"
st.markdown(f"""
<style>
.stApp {{background:#f5f7fb;color:#172b4d}}
[data-testid="stSidebar"] {{background:{NAVY};}}
[data-testid="stSidebar"] * {{color:white !important;}}
h1,h2,h3 {{color:{NAVY};}}
.pms-hero {{background:linear-gradient(120deg,{NAVY},#24517d);padding:24px 28px;border-radius:16px;color:white;margin-bottom:16px}}
.pms-hero h1 {{color:white;margin:0}}
.pms-hero p {{color:#f6df9b;margin:8px 0 0 0}}
div.stButton>button {{border-radius:9px;border:1px solid {GOLD};}}
</style>
""", unsafe_allow_html=True)

def init_state():
    defaults = {
        "tasks": pd.DataFrame([
            {"ID":"T-001","Task":"Define project charter","Owner":"Project Manager","Sprint":"Sprint 1","Status":"In Progress","Priority":"High","Estimate (hrs)":8,"Budget (PKR)":20000,"Due date":str(datetime.date.today()+datetime.timedelta(days=5))},
            {"ID":"T-002","Task":"Stakeholder requirements","Owner":"Business Analyst","Sprint":"Sprint 1","Status":"To Do","Priority":"High","Estimate (hrs)":12,"Budget (PKR)":30000,"Due date":str(datetime.date.today()+datetime.timedelta(days=8))},
            {"ID":"T-003","Task":"Create initial prototype","Owner":"Development Team","Sprint":"Sprint 2","Status":"To Do","Priority":"Medium","Estimate (hrs)":20,"Budget (PKR)":60000,"Due date":str(datetime.date.today()+datetime.timedelta(days=15))}
        ]),
        "risks": pd.DataFrame([{"Risk":"Scope creep","Probability":"Medium","Impact":"High","Mitigation":"Baseline scope; review changes through backlog refinement","Owner":"Project Manager","Status":"Open"}]),
        "team": pd.DataFrame([
            {"Role":"Product Owner","Responsibilities":"Owns product vision, value, and backlog priority","Accountability":"Maximize product value"},
            {"Role":"Scrum Master / Agile Coach","Responsibilities":"Facilitates ceremonies and removes impediments","Accountability":"Enable effective Agile delivery"},
            {"Role":"Project Manager","Responsibilities":"Budget, dependencies, governance, reporting and stakeholder coordination","Accountability":"Delivery within agreed constraints"},
            {"Role":"Business Analyst","Responsibilities":"Elicit requirements; define acceptance criteria","Accountability":"Clear, testable requirements"},
            {"Role":"Developers / Delivery Team","Responsibilities":"Design, build, test and integrate increments","Accountability":"Quality working increment"},
            {"Role":"QA / Test Engineer","Responsibilities":"Test strategy, automation, defect verification","Accountability":"Quality evidence"},
            {"Role":"Finance / Budget Owner","Responsibilities":"Cost baseline, forecasts, approvals and variance review","Accountability":"Financial control"},
            {"Role":"Stakeholders / Sponsor","Responsibilities":"Set outcomes, approve funding and key decisions","Accountability":"Business alignment"}
        ])
    }
    for k,v in defaults.items():
        if k not in st.session_state: st.session_state[k]=v
init_state()

st.sidebar.markdown("## 📊 PMS")
st.sidebar.caption("Project Management System")
page = st.sidebar.radio("Workspace", ["Executive Dashboard","Project Setup","Budget & Allocation","Agile Workboard","Team & RACI","Risks & Issues","AI Project Advisor","Reports & Downloads","Settings"])
st.sidebar.markdown("---")
st.sidebar.markdown("**Designed by SkyWood Innovator**")
st.sidebar.caption("Plan clearly. Allocate wisely. Deliver value.")

if "project" not in st.session_state:
    st.session_state.project = {"name":"PMS Demo Project","description":"A sample project workspace","sponsor":"Project Sponsor","manager":"Project Manager","method":"Scrum","currency":"PKR","budget":1000000,"contingency":10,"start":datetime.date.today(),"end":datetime.date.today()+datetime.timedelta(days=90),"goal":"Deliver measurable business value with transparent budget and quality controls."}
p=st.session_state.project
tasks=st.session_state.tasks

def money(x):
    try: return f"{p.get('currency','PKR')} {float(x):,.0f}"
    except: return str(x)

def build_report_text():
    total_budget=float(p.get("budget",0))
    allocated=float(pd.to_numeric(st.session_state.tasks["Budget (PKR)"], errors="coerce").fillna(0).sum()) if len(st.session_state.tasks) else 0
    completed=(st.session_state.tasks["Status"]=="Done").sum() if len(st.session_state.tasks) else 0
    return f"""PMS PROJECT MANAGEMENT REPORT
Project: {p['name']}
Description: {p['description']}
Sponsor: {p['sponsor']}
Project Manager: {p['manager']}
Framework: {p['method']}
Period: {p['start']} to {p['end']}
Currency: {p['currency']}
Approved budget: {money(total_budget)}
Task-level allocation: {money(allocated)}
Unallocated budget: {money(total_budget-allocated)}
Tasks: {len(st.session_state.tasks)} | Completed: {completed}
Business goal: {p['goal']}

BUDGET ALLOCATION
{st.session_state.allocations.to_string(index=False) if "allocations" in st.session_state and len(st.session_state.allocations) else "No category allocation entered."}

TASK REGISTER
{st.session_state.tasks.to_string(index=False)}

TEAM ROLES
{st.session_state.team.to_string(index=False)}

RISKS AND ISSUES
{st.session_state.risks.to_string(index=False)}

AGILE WORKFLOW
1. Discover: define problem, stakeholders, outcomes, constraints and success measures.
2. Initiate: create project charter, product vision, high-level estimate and budget baseline.
3. Plan: build prioritized product backlog, definition of done, release plan, dependencies and risk register.
4. Sprint planning: select work based on capacity, priority and sprint goal.
5. Deliver: daily coordination, design/build/test, continuous integration and quality checks.
6. Review: demonstrate a potentially releasable increment and collect stakeholder feedback.
7. Retrospective: identify improvements and assign owners.
8. Monitor: review budget burn, forecast at completion, scope, risks, quality and benefits.
9. Release/close: acceptance, handover, lessons learned and financial reconciliation.

Budget controls: approve a cost baseline; assign each work package an owner; forecast actuals and estimate-to-complete; investigate material variance; use change control for scope or funding changes. Contingency is a planning reserve, not automatic spending authority.

Governance note: adapt roles and ceremonies to project size. Scrum roles/events and Kanban flow practices can be combined thoughtfully, but teams should avoid mixing frameworks without clear working agreements.
"""
def make_docx():
    d=Document(); d.add_heading("PMS Project Management Report",0)
    for line in build_report_text().splitlines():
        if line.strip().isupper() and len(line)<50: d.add_heading(line.title(), level=1)
        elif line.strip(): d.add_paragraph(line)
    b=io.BytesIO(); d.save(b); return b.getvalue()
def make_pdf():
    b=io.BytesIO(); doc=SimpleDocTemplate(b,pagesize=A4,rightMargin=38,leftMargin=38,topMargin=38,bottomMargin=38)
    styles=getSampleStyleSheet(); story=[]
    for line in build_report_text().splitlines():
        if not line.strip(): story.append(Spacer(1,6))
        elif line.strip().isupper() and len(line)<50: story.append(Paragraph(f"<b>{line.title()}</b>",styles["Heading2"]))
        else: story.append(Paragraph(line.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;"),styles["BodyText"]))
        story.append(Spacer(1,3))
    doc.build(story); return b.getvalue()

st.markdown(f'<div class="pms-hero"><h1>📊 PMS — Project Management System</h1><p>Agile delivery • Smart budget allocation • AI-assisted decisions • Executive-ready reports</p><p><b>Designed by SkyWood Innovator</b></p></div>',unsafe_allow_html=True)

if page=="Executive Dashboard":
    st.subheader(p["name"])
    st.caption(p["description"])
    task_budget=pd.to_numeric(tasks["Budget (PKR)"],errors="coerce").fillna(0).sum() if len(tasks) else 0
    done=(tasks["Status"]=="Done").sum() if len(tasks) else 0
    c1,c2,c3,c4=st.columns(4)
    c1.metric("Approved Budget",money(p["budget"]))
    c2.metric("Task Allocation",money(task_budget),f"{(task_budget/max(float(p['budget']),1))*100:.1f}% of baseline")
    c3.metric("Unallocated",money(float(p["budget"])-task_budget))
    c4.metric("Tasks Completed",f"{done}/{len(tasks)}")
    left,right=st.columns(2)
    with left:
        st.markdown("### Task status")
        if len(tasks): st.bar_chart(tasks["Status"].value_counts())
        else: st.info("Add tasks in Agile Workboard.")
    with right:
        st.markdown("### Budget by work item")
        if len(tasks): st.bar_chart(tasks.set_index("Task")["Budget (PKR)"])
        else: st.info("No budgeted tasks yet.")
    st.markdown("### Delivery health checks")
    for label,ok in [
        ("Project dates defined",bool(p.get("start") and p.get("end"))),
        ("Task allocation within approved budget",task_budget<=float(p["budget"])),
        ("At least one owner assigned",len(tasks)>0 and (tasks["Owner"].astype(str).str.strip()!="").all()),
        ("Risks being tracked",len(st.session_state.risks)>0)
    ]:
        st.write(("✅ " if ok else "⚠️ ")+label)

elif page=="Project Setup":
    st.subheader("Project charter & baseline")
    with st.form("project_form"):
        name=st.text_input("Project name",p["name"]); desc=st.text_area("Description",p["description"])
        a,b=st.columns(2)
        sponsor=a.text_input("Sponsor / client",p["sponsor"]); manager=b.text_input("Project manager",p["manager"])
        method=st.selectbox("Delivery framework",["Scrum","Kanban","Scrumban","Hybrid Agile","Predictive / Waterfall"],index=["Scrum","Kanban","Scrumban","Hybrid Agile","Predictive / Waterfall"].index(p["method"]))
        currency=st.selectbox("Budget currency",["PKR","USD","GBP","EUR","AED"],index=["PKR","USD","GBP","EUR","AED"].index(p["currency"]))
        budget=st.number_input("Approved budget",min_value=0.0,value=float(p["budget"]),step=10000.0)
        contingency=st.slider("Contingency reserve (%)",0,30,int(p["contingency"]))
        a,b=st.columns(2); start=a.date_input("Start date",p["start"]); end=b.date_input("Target end date",p["end"])
        goal=st.text_area("Business goal / success outcome",p["goal"])
        submit=st.form_submit_button("Save project baseline",type="primary")
        if submit:
            if end<start: st.error("Target end date must be on or after the start date.")
            else:
                st.session_state.project={"name":name,"description":desc,"sponsor":sponsor,"manager":manager,"method":method,"currency":currency,"budget":budget,"contingency":contingency,"start":start,"end":end,"goal":goal}
                st.success("Project baseline saved for this session.")

elif page=="Budget & Allocation":
    st.subheader("Budget planning and allocation")
    st.write("Use category-level allocations to create a transparent cost baseline. Keep contingency separate from committed spend.")
    categories=["People & Labour","Software & Cloud","Equipment & Materials","Training","Travel & Logistics","Testing & Quality","Marketing & Communications","Operations","Contingency Reserve","Other"]
    if "allocations" not in st.session_state:
        st.session_state.allocations=pd.DataFrame([{"Category":c,"Allocation":0.0,"Owner":"","Notes":""} for c in categories])
        st.session_state.allocations.loc[0,"Allocation"]=float(p["budget"])*0.45
        st.session_state.allocations.loc[1,"Allocation"]=float(p["budget"])*0.15
        st.session_state.allocations.loc[2,"Allocation"]=float(p["budget"])*0.15
        st.session_state.allocations.loc[8,"Allocation"]=float(p["budget"])*p["contingency"]/100
    st.caption(f"Approved baseline: {money(p['budget'])} | Suggested contingency: {p['contingency']}%")
    st.session_state.allocations=st.data_editor(st.session_state.allocations,num_rows="dynamic",use_container_width=True,key="alloc_editor",column_config={"Allocation":st.column_config.NumberColumn(min_value=0,step=1000,format="%.2f")})
    total=float(pd.to_numeric(st.session_state.allocations["Allocation"],errors="coerce").fillna(0).sum())
    c1,c2,c3=st.columns(3); c1.metric("Allocated",money(total)); c2.metric("Unallocated / (Over)",money(float(p["budget"])-total)); c3.metric("Allocation ratio",f"{total/max(float(p['budget']),1)*100:.1f}%")
    if total>float(p["budget"]): st.error("Allocation exceeds approved budget. Reduce allocations or request formal approval for a revised baseline.")
    elif total<float(p["budget"]): st.info("Some budget remains unallocated. Consider leaving it unassigned until estimates are clearer.")
    st.download_button("Download allocation CSV",st.session_state.allocations.to_csv(index=False).encode("utf-8"),"pms_budget_allocation.csv","text/csv")
    st.markdown("### Budget allocation principles")
    st.markdown("- Estimate from deliverables and work packages, not arbitrary percentages alone.\n- Separate baseline, committed cost, actual cost, forecast and contingency.\n- Assign an owner and approval threshold to each cost category.\n- Review variance each sprint or reporting period; document changes.\n- Forecast Estimate at Completion (EAC) = Actual Cost (AC) + Estimate to Complete (ETC).")

elif page=="Agile Workboard":
    st.subheader("Agile backlog & delivery workboard")
    st.caption("The board is a lightweight work register; adapt sprint length, WIP limits and ceremonies to your team.")
    edited=st.data_editor(st.session_state.tasks,num_rows="dynamic",use_container_width=True,key="tasks_editor",column_config={"Status":st.column_config.SelectboxColumn(options=["Backlog","To Do","In Progress","Review","Blocked","Done"]),"Priority":st.column_config.SelectboxColumn(options=["Low","Medium","High","Critical"]),"Budget (PKR)":st.column_config.NumberColumn(min_value=0,step=1000),"Estimate (hrs)":st.column_config.NumberColumn(min_value=0,step=1)})
    st.session_state.tasks=edited
    st.download_button("Export task register CSV",edited.to_csv(index=False).encode("utf-8"),"pms_task_register.csv","text/csv")
    st.markdown("### Suggested Scrum cadence")
    st.markdown("- **Sprint Planning:** agree on a Sprint Goal and select work based on priority and capacity.\n- **Daily Scrum:** inspect progress toward the goal and adapt the plan.\n- **Backlog Refinement:** clarify upcoming items and acceptance criteria.\n- **Sprint Review:** demonstrate the increment and gather feedback.\n- **Retrospective:** select one or two measurable process improvements.")
    st.warning("A status board does not replace a shared Definition of Done, tested increments, or stakeholder acceptance.")

elif page=="Team & RACI":
    st.subheader("Team roles and responsibility model")
    st.session_state.team=st.data_editor(st.session_state.team,num_rows="dynamic",use_container_width=True,key="team_editor")
    st.download_button("Download team roles CSV",st.session_state.team.to_csv(index=False).encode("utf-8"),"pms_team_roles.csv","text/csv")
    st.markdown("### Example RACI matrix")
    raci=pd.DataFrame([
        {"Activity":"Product vision & outcomes","Sponsor":"A","Product Owner":"R","PM":"C","Team":"I","Finance":"C"},
        {"Activity":"Prioritize backlog","Sponsor":"I","Product Owner":"A/R","PM":"C","Team":"C","Finance":"I"},
        {"Activity":"Sprint delivery & quality","Sponsor":"I","Product Owner":"C","PM":"C","Team":"A/R","Finance":"I"},
        {"Activity":"Budget baseline & forecast","Sponsor":"A","Product Owner":"C","PM":"R","Team":"C","Finance":"R"},
        {"Activity":"Risk and dependency review","Sponsor":"I","Product Owner":"C","PM":"A/R","Team":"R","Finance":"C"},
        {"Activity":"Release acceptance","Sponsor":"A","Product Owner":"R","PM":"C","Team":"R","Finance":"I"}
    ])
    st.dataframe(raci,use_container_width=True,hide_index=True)
    st.caption("R = Responsible, A = Accountable, C = Consulted, I = Informed. Confirm one accountable owner per decision in your governance charter.")

elif page=="Risks & Issues":
    st.subheader("Risk, issue and decision register")
    st.session_state.risks=st.data_editor(st.session_state.risks,num_rows="dynamic",use_container_width=True,key="risks_editor",column_config={"Probability":st.column_config.SelectboxColumn(options=["Low","Medium","High"]),"Impact":st.column_config.SelectboxColumn(options=["Low","Medium","High","Critical"]),"Status":st.column_config.SelectboxColumn(options=["Open","Mitigating","Escalated","Closed"])})
    st.download_button("Download risk register CSV",st.session_state.risks.to_csv(index=False).encode("utf-8"),"pms_risk_register.csv","text/csv")
    st.markdown("Review risks at planning and regularly during delivery. Each risk should have a named owner, a trigger, a mitigation, and a review date.")

elif page=="AI Project Advisor":
    st.subheader("OpenAI-powered project advisor")
    st.write("Ask for project-specific help with budget allocation, agile planning, risk analysis, sprint goals, role design, or executive summaries.")
    api_key=st.text_input("OpenAI API key",type="password",value=os.getenv("OPENAI_API_KEY",""),help="For deployment, store OPENAI_API_KEY in Streamlit Secrets instead of hard-coding it.")
    model=st.text_input("OpenAI model",value=os.getenv("OPENAI_MODEL","gpt-4.1-mini"))
    question=st.text_area("What should the AI help with?",placeholder="Example: Review my budget allocation and suggest a balanced allocation with assumptions.")
    if st.button("Ask PMS AI",type="primary"):
        if not api_key: st.error("Add your OpenAI API key or configure OPENAI_API_KEY in Streamlit Secrets.")
        elif not question.strip(): st.warning("Enter a question first.")
        else:
            try:
                client=OpenAI(api_key=api_key)
                context={"project":p,"tasks":st.session_state.tasks.fillna("").to_dict(orient="records"),"allocations":st.session_state.get("allocations",pd.DataFrame()).fillna("").to_dict(orient="records"),"risks":st.session_state.risks.fillna("").to_dict(orient="records"),"team":st.session_state.team.fillna("").to_dict(orient="records")}
                with st.spinner("Reviewing your project context..."):
                    response=client.responses.create(model=model,instructions="You are PMS, a careful senior project manager and Agile delivery advisor. Provide practical recommendations, assumptions, formulas and next actions. Never claim to approve spending or guarantee estimates. Flag missing data and financial risks. Use Scrum Guide concepts accurately and tailor the framework to project context.",input="Project context:\n"+json.dumps(context,default=str)+"\n\nUser request:\n"+question)
                st.markdown(response.output_text)
            except Exception as e: st.error(f"AI request failed: {e}")

elif page=="Reports & Downloads":
    st.subheader("Project reports & exports")
    st.write("Generate a snapshot report from your current project, task register, team roles, risk register, and budget allocations.")
    st.text_area("Report preview",build_report_text(),height=300)
    a,b,c=st.columns(3)
    a.download_button("Download Word report (.docx)",make_docx(),"PMS_Project_Report.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)
    b.download_button("Download PDF report",make_pdf(),"PMS_Project_Report.pdf","application/pdf",use_container_width=True)
    c.download_button("Download report (.txt)",build_report_text(),"PMS_Project_Report.txt","text/plain",use_container_width=True)
    st.download_button("Download full project data (.json)",json.dumps({"project":p,"tasks":st.session_state.tasks.fillna("").to_dict(orient="records"),"allocations":st.session_state.get("allocations",pd.DataFrame()).fillna("").to_dict(orient="records"),"team":st.session_state.team.fillna("").to_dict(orient="records"),"risks":st.session_state.risks.fillna("").to_dict(orient="records")},default=str,indent=2),"PMS_Project_Data.json","application/json")

elif page=="Settings":
    st.subheader("Configuration & deployment")
    st.markdown("### OpenAI API configuration")
    st.code('OPENAI_API_KEY = "your-api-key"\nOPENAI_MODEL = "gpt-4.1-mini"',language="toml")
    st.write("Place these values in `.streamlit/secrets.toml` locally or in your hosting provider's Streamlit Secrets. Do not commit API keys to GitHub.")
    st.markdown("### Data persistence")
    st.warning("This starter version keeps edits in Streamlit session state. Refreshing the app or ending the session may reset data. For production or multi-user use, add authenticated access and a persistent database such as PostgreSQL or SQLite with backups.")
    st.markdown("### Agile operating workflow")
    st.markdown("**Discover → Charter & Fund → Prioritize Backlog → Plan Sprint/Flow → Build & Test → Review/Accept → Retrospective → Forecast & Adapt → Release & Close**")
    st.markdown("PMS is a planning aid, not a substitute for finance approval, procurement controls, legal review, or formal project governance.")

st.markdown("---")
st.caption("PMS • Designed by SkyWood Innovator • Budget decisions remain subject to authorized human review.")
